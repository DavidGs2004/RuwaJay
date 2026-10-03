import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build_engineering_document():
    doc = docx.Document()

    # Configuración de márgenes estándar (1 pulgada / 2.54 cm)
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    COLOR_TERRACOTA = RGBColor(234, 88, 12)  # #EA580C
    COLOR_CAFE = RGBColor(30, 41, 59)        # #1E293B
    COLOR_MUTED = RGBColor(100, 116, 139)     # #64748B

    def add_custom_heading(text, level, color):
        h = doc.add_heading(level=level)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.color.rgb = color
        return h

    def set_cell_style(cell, bg_color="FFFFFF", top=90, bottom=90, left=130, right=130):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        tcPr.append(shd)
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

    def draw_diagram_card(steps, title_text):
        add_custom_heading(title_text, level=2, color=COLOR_CAFE)

        tbl = doc.add_table(rows=0, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        for i, (kind, text, detail) in enumerate(steps):
            row = tbl.add_row()
            c = row.cells[0]
            
            if kind == "start_end":
                bg = "FED7AA"     # Naranja suave
                prefix = "● [INICIO/FIN] "
            elif kind == "decision":
                bg = "FEF3C7"     # Amarillo
                prefix = "◆ [DECISIÓN] "
            elif kind == "db":
                bg = "E0E7FF"     # Índigo suave
                prefix = "🖴 [FIRESTORE / STORAGE] "
            else:
                bg = "F1F5F9"     # Gris slate
                prefix = "■ [OPERACIÓN] "

            set_cell_style(c, bg_color=bg, top=90, bottom=90, left=150, right=150)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            rt = p.add_run(f"{prefix}{text}\n")
            rt.font.name = "Arial"
            rt.font.size = Pt(9.5)
            rt.font.bold = True
            rt.font.color.rgb = COLOR_CAFE

            rd = p.add_run(detail)
            rd.font.name = "Arial"
            rd.font.size = Pt(8.5)
            rd.font.color.rgb = COLOR_MUTED

            if i < len(steps) - 1:
                r_arr = tbl.add_row()
                c_arr = r_arr.cells[0]
                set_cell_style(c_arr, bg_color="FFFFFF", top=25, bottom=25, left=150, right=150)
                p_arr = c_arr.paragraphs[0]
                p_arr.alignment = WD_ALIGN_PARAGRAPH.CENTER
                arr_run = p_arr.add_run("▼")
                arr_run.font.name = "Arial"
                arr_run.font.size = Pt(11)
                arr_run.font.bold = True
                arr_run.font.color.rgb = COLOR_TERRACOTA

        doc.add_paragraph()

    # ==========================
    # 1. PORTADA FORMAL
    # ==========================
    p_pre = doc.add_paragraph()
    p_pre.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_pre = p_pre.add_run("INGENIERÍA EN SISTEMAS / INFORMÁTICA • 8VO SEMESTRE\nDESARROLLO DE SISTEMAS WEB Y APLICACIONES MÓVILES")
    r_pre.font.name = "Arial"
    r_pre.font.size = Pt(9)
    r_pre.font.bold = True
    r_pre.font.color.rgb = COLOR_MUTED

    doc.add_paragraph("\n\n")

    p_t = doc.add_paragraph()
    r_t = p_t.add_run("DOCUMENTACIÓN TÉCNICA DE INGENIERÍA DE SOFTWARE\nSISTEMA INMOBILIARIO RUWAJAY")
    r_t.font.name = "Arial"
    r_t.font.size = Pt(22)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_TERRACOTA

    p_s = doc.add_paragraph()
    r_s = p_s.add_run("Especificación de Requerimientos del Sistema, Arquitectura React + Jetpack Compose, Modelado NoSQL, Trazabilidad y Casos de Prueba")
    r_s.font.name = "Arial"
    r_s.font.size = Pt(12)
    r_s.font.color.rgb = COLOR_CAFE

    doc.add_paragraph("\n" * 5)

    p_m = doc.add_paragraph()
    r_m = p_m.add_run(
        "Proyecto: RuwaJay Platform Multiplataforma\n"
        "Infraestructura: Serverless BaaS (Firebase Cloud + Hosting CDN + OSRM Engine)\n"
        "Fecha de Emisión: Octubre 2026\n"
        "Estado del Entregable: Versión Final del Proyecto"
    )
    r_m.font.name = "Arial"
    r_m.font.size = Pt(10)
    r_m.font.color.rgb = COLOR_MUTED

    doc.add_page_break()

    # ==========================
    # 2. INTRODUCCIÓN Y METODOLOGÍA
    # ==========================
    add_custom_heading("1. Introducción y Marco de Trabajo", level=1, color=COLOR_TERRACOTA)
    
    doc.add_paragraph(
        "RuwaJay es una plataforma integral de software (SPA Web y Aplicación Nativa Android) diseñada para optimizar "
        "y facilitar el arrendamiento y compra de vivienda, conectando a buscadores (seekers) con propietarios (owners). "
        "El sistema integra exploración geoespacial interactiva, cálculo algorítmico de rutas viales (OSRM Engine), matriz comparativa técnica, "
        "mensajería instantánea transaccional y control estricto de sesiones bajo una arquitectura en la nube de alta disponibilidad."
    )
    
    doc.add_paragraph(
        "Metodología de Desarrollo: El proyecto se estructuró mediante metodología ágil con entregas iterativas y control de versiones bajo Git Flow, "
        "asegurando la sincronización simultánea del cliente web en React 18 y el cliente móvil en Kotlin con Jetpack Compose."
    )

    # ==========================
    # 3. REQUERIMIENTOS FUNCIONALES
    # ==========================
    add_custom_heading("2. Especificación de Requerimientos Funcionales (RF)", level=1, color=COLOR_TERRACOTA)

    rf_data = [
        ["RF-01", "Autenticación de Usuarios y Roles", "El sistema debe autenticar mediante Correo/Contraseña y Google Sign-In, asignando roles diferenciados: Seeker (Buscador) y Owner (Propietario).", "Alta"],
        ["RF-02", "Control y Cierre de Sesión por Inactividad", "El sistema debe detectar inactividad en la estación de trabajo y desplegar un modal de advertencia a los 5 minutos con auto-logout en 60 segundos.", "Alta"],
        ["RF-03", "Restablecimiento Seguro de Contraseña", "El sistema debe permitir la recuperación de acceso capturando el token 'oobCode' en la ruta interna /reset-password sin usar pantallas externas.", "Media"],
        ["RF-04", "Publicación de Viviendas", "Debe permitir exclusivamente a usuarios con rol Owner (Propietario) subir descripciones, especificaciones, ubicación y fotos a Storage.", "Alta"],
        ["RF-05", "Gestión de Propiedades Favoritas", "El usuario debe poder marcar y desmarcar inmuebles favoritos con sincronización y persistencia inmediata en Web y Android.", "Media"],
        ["RF-06", "Comparador Multi-inmueble", "Debe permitir seleccionar hasta 3 propiedades en simultáneo y desplegar una matriz técnica comparativa de características.", "Media"],
        ["RF-07", "Mensajería en Tiempo Real", "Debe habilitar una sala de chat bidireccional inmediata una vez que el propietario confirme la solicitud de visita a un inmueble.", "Alta"],
        ["RF-08", "Cálculo de Ruta Vial y Enlace a Waze/Maps", "Debe calcular la ruta vial óptima mediante OSRM Engine, estimar tiempos y permitir deep linking directo a las apps de Waze y Google Maps.", "Alta"],
        ["RF-09", "Sincronización Reactiva Multiplataforma", "El sistema debe actualizar automáticamente los cambios en Web y Android vía WebSockets sin necesidad de recargar la página.", "Alta"],
        ["RF-10", "Avisos Globales del Sistema", "Debe consumir y desplegar avisos técnicos o de mantenimiento sincronizados desde la colección system_updates.", "Baja"]
    ]

    tbl_rf = doc.add_table(rows=1, cols=4)
    tbl_rf.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_rf = tbl_rf.rows[0].cells
    for i, h in enumerate(["ID", "Requerimiento", "Descripción Funcional", "Prioridad"]):
        c_rf[i].text = h
        set_cell_style(c_rf[i], bg_color="EA580C")
        c_rf[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c_rf[i].paragraphs[0].runs[0].font.bold = True
        c_rf[i].paragraphs[0].runs[0].font.name = "Arial"
        c_rf[i].paragraphs[0].runs[0].font.size = Pt(9)

    for row in rf_data:
        r_c = tbl_rf.add_row().cells
        for idx, val in enumerate(row):
            r_c[idx].text = val
            set_cell_style(r_c[idx], bg_color="FFFFFF" if idx != 0 else "F8FAFC")
            r_c[idx].paragraphs[0].runs[0].font.name = "Arial"
            r_c[idx].paragraphs[0].runs[0].font.size = Pt(8.5)

    # ==========================
    # 4. ARQUITECTURA Y RBAC
    # ==========================
    doc.add_page_break()
    add_custom_heading("3. Arquitectura del Sistema y Control de Acceso (RBAC)", level=1, color=COLOR_TERRACOTA)

    doc.add_paragraph(
        "El proyecto opera bajo una arquitectura Serverless BaaS donde la lógica de presentación reside en las aplicaciones web y móviles, "
        "y la persistencia y seguridad es gestionada por Firebase Cloud y servicios públicos:"
    )
    doc.add_paragraph("• Frontend Web: React 18 + Vite + TailwindCSS, orquestado mediante Context Providers modulares.", style='List Bullet')
    doc.add_paragraph("• Frontend Móvil: Kotlin nativo con Jetpack Compose y arquitectura MVVM + Repository Pattern.", style='List Bullet')
    doc.add_paragraph("• Capa de Persistencia: Cloud Firestore con sincronización en tiempo real vía WebSockets.", style='List Bullet')
    doc.add_paragraph("• Almacenamiento de Medios: Firebase Storage para imágenes de propiedades.", style='List Bullet')
    doc.add_paragraph("• Georuteo Vial: OSRM Engine sobre OpenStreetMap integrado con Leaflet.js y osmdroid.", style='List Bullet')

    add_custom_heading("3.1. Matriz de Control de Acceso Basado en Roles (RBAC)", level=2, color=COLOR_CAFE)
    rbac_data = [
        ["Explorar catálogo y mapa", "Permitido", "Permitido", "Permitido", "Permitido"],
        ["Comparar hasta 3 inmuebles", "Permitido", "Permitido", "Permitido", "Permitido"],
        ["Gestionar favoritos", "Denegado", "Permitido", "Permitido", "Permitido"],
        ["Solicitar cita / visita", "Denegado", "Permitido", "Denegado", "Permitido"],
        ["Confirmar visitas recibidas", "Denegado", "Denegado", "Permitido", "Permitido"],
        ["Chatear tras confirmación", "Denegado", "Permitido", "Permitido", "Permitido"],
        ["Publicar / Modificar inmuebles", "Denegado", "Denegado", "Permitido", "Permitido"],
        ["Emitir system_updates", "Denegado", "Denegado", "Denegado", "Permitido (Admin)"]
    ]

    tbl_rbac = doc.add_table(rows=1, cols=5)
    tbl_rbac.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_rb = tbl_rbac.rows[0].cells
    for i, h in enumerate(["Operación / Módulo", "Anónimo", "Seeker (Buscador)", "Owner (Propietario)", "Administrador"]):
        c_rb[i].text = h
        set_cell_style(c_rb[i], bg_color="EA580C")
        c_rb[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c_rb[i].paragraphs[0].runs[0].font.bold = True
        c_rb[i].paragraphs[0].runs[0].font.name = "Arial"
        c_rb[i].paragraphs[0].runs[0].font.size = Pt(8.5)

    for row in rbac_data:
        r_c = tbl_rbac.add_row().cells
        for idx, val in enumerate(row):
            r_c[idx].text = val
            bg = "FEF2F2" if val == "Denegado" else ("F0FDF4" if val.startswith("Permitido") else "F8FAFC")
            set_cell_style(r_c[idx], bg_color=bg)
            r_c[idx].paragraphs[0].runs[0].font.name = "Arial"
            r_c[idx].paragraphs[0].runs[0].font.size = Pt(8)

    # ==========================
    # 5. DIAGRAMAS DE FLUJO
    # ==========================
    doc.add_page_break()
    add_custom_heading("4. Diagramas de Flujo de Procesos Clave", level=1, color=COLOR_TERRACOTA)

    # 4.1 Autenticación
    draw_diagram_card([
        ("start_end", "Inicio de App o Web", "El cliente abre el aplicativo en el navegador o dispositivo móvil"),
        ("decision", "¿Hay sesión activa en el dispositivo?", "Verificación local de credenciales / tokens en caché"),
        ("process", "Pantalla de Login", "Despliegue de opciones: Correo/Contraseña, Google Sign-In o Registro"),
        ("decision", "¿Qué acción elige el usuario?", "Bifurcación según la acción ingresada por el cliente"),
        ("process", "Ingreso de Credenciales o Google Auth", "Ejecución de signInWithEmailAndPassword o GoogleProvider"),
        ("db", "Firebase Auth valida y consulta Firestore 'users'", "Se verifica la identidad y se extrae el rol: Seeker o Owner"),
        ("start_end", "Redirigir a Pantalla Principal / Home", "Ingreso concedido según permisos del perfil")
    ], "4.1. Flujo de Autenticación (Login / Registro / Roles)")

    # 4.2 Inactividad
    draw_diagram_card([
        ("start_end", "Sesión de Usuario Activa", "El usuario navega dentro del portal autenticado"),
        ("process", "Listener de Interacciones DOM", "Monitoreo continuo de eventos: clics, teclado, scroll y mouse"),
        ("decision", "¿Tiempo inactivo >= 5 minutos?", "Validación contra ruwajay_last_activity en LocalStorage"),
        ("process", "Despliegue de Modal: RuwaJay Security", "Ventana emergente con cuenta regresiva de 60 segundos"),
        ("decision", "¿El usuario pulsa 'Mantener Sesión'?", "Confirmación de presencia física frente a la pantalla"),
        ("db", "Ejecutar signOut en Firebase Auth", "Cierre preventivo de credenciales por expiración de tiempo"),
        ("start_end", "Redirección Forzosa a /login", "Protección contra Session Hijacking y sesión desatendida")
    ], "4.2. Flujo de Monitor de Inactividad y Caducidad de Sesión")

    # 4.3 Restablecer Clave
    draw_diagram_card([
        ("start_end", "Clic en '¿Olvidaste tu contraseña?'", "El usuario solicita auxilio de credenciales desde el Login"),
        ("process", "Ingreso de Correo Electrónico", "Captura de email y llamada a requestPasswordReset"),
        ("process", "Recepción de Correo y Clic en Link", "El usuario abre el link seguro que contiene el token 'oobCode'"),
        ("process", "Intercepción en ResetPasswordPage", "La ruta /reset-password captura el código seguro de la URL"),
        ("db", "Firebase verifyPasswordResetCode", "Verificación criptográfica de validez y caducidad"),
        ("process", "Ingreso de Nueva Contraseña", "Validación en cliente: mínimo 6 caracteres y coincidencia exacta"),
        ("db", "Firebase confirmPasswordReset", "Actualización efectiva de credenciales en el proveedor Auth"),
        ("start_end", "Redirección Automática al Login", "Acceso rehabilitado con la nueva contraseña")
    ], "4.3. Flujo de Restablecimiento Seguro de Contraseña")

    # 4.4 Publicación
    draw_diagram_card([
        ("start_end", "Acceso a Pantalla de Publicación", "Ruta /publicar en Web o formulario de alta en Android"),
        ("decision", "¿El usuario tiene rol Owner (Propietario)?", "Comprobación del atributo role en el documento del usuario"),
        ("process", "Llenar Formulario de Propiedad", "Captura de título, descripción, precio, especificaciones y dirección"),
        ("db", "Subir Imágenes a Firebase Storage", "Carga asíncrona de archivos binarios en el bucket 'properties/'"),
        ("process", "Obtener URLs Públicas", "Recuperación de los enlaces permanentes de descarga"),
        ("db", "Crear Documento en Firestore 'properties'", "Escritura atómica del inmueble con geolocalización lat/lng"),
        ("start_end", "Propiedad Visible en Explore / Home", "Propagación automática e inmediata a todos los clientes")
    ], "4.4. Flujo de Publicación de Propiedades (Rol Propietario)")

    # 4.5 Favoritos
    draw_diagram_card([
        ("start_end", "Interacción con Card de Propiedad", "El usuario hace clic sobre el icono de corazón"),
        ("decision", "¿Inmueble ya marcado como favorito?", "Comprobación en el arreglo de favoritos de FavoritesContext"),
        ("db", "Firestore: Actualizar Array de Favoritos", "Ejecución de arrayUnion o arrayRemove sobre 'users/UID/favorites'"),
        ("process", "Actualización de Estado Visual", "Cambio cromático del botón de corazón sin recargar pantalla"),
        ("start_end", "Sincronización en Pestaña Favoritos", "Visualización actualizada en ProfilePage y ProfilesScreen.kt")
    ], "4.5. Flujo de Gestión de Favoritos en Tiempo Real")

    # 4.6 Comparador
    draw_diagram_card([
        ("start_end", "Selección de Opción 'Comparar'", "El usuario pulsa el selector en una propiedad del catálogo"),
        ("decision", "¿Existen ya 3 propiedades en comparación?", "Verificación del límite ergonómico de comparación"),
        ("process", "Actualizar Matriz en CompareContext", "Almacenamiento temporal de las propiedades a cotejar"),
        ("process", "Apertura de PropertyCompareModal / Dialog", "Despliegue de la tabla técnica multi-columna"),
        ("start_end", "Análisis Comparativo Paramétrico", "Inspección lado a lado de precio, metraje, habitaciones y zona")
    ], "4.6. Flujo de Comparación de Propiedades")

    # 4.7 Chat
    draw_diagram_card([
        ("start_end", "Seeker (Buscador) Solicita Visita a Propiedad", "Envío de solicitud formal desde la ficha del inmueble"),
        ("process", "Owner (Propietario) Confirma la Visita", "Aceptación de la cita mediante el panel de control"),
        ("db", "Se Habilita Sala en Firestore 'conversations'", "Creación o activación del canal vinculado a la propiedad"),
        ("process", "Envío de Mensajes", "Escritura atómica en la subcolección 'messages' con timestamp de servidor"),
        ("db", "Firestore onSnapshot Listener", "Detección instantánea de nuevos mensajes mediante WebSockets"),
        ("start_end", "Actualización Inmediata de UI", "Renderizado reactivo simultáneo en el cliente Web y Android")
    ], "4.7. Flujo del Chat en Tiempo Real y Coordinación")

    # 4.8 Navegación
    draw_diagram_card([
        ("start_end", "Usuario Solicita Ruta hacia Inmueble", "Petición 'Cómo llegar' desde la pantalla de detalle"),
        ("process", "Captura de Coordenadas GPS y Destino", "Geolocalización del cliente y lat/lng de la propiedad"),
        ("process", "Consulta al Motor OSRM Engine", "Petición HTTP a la API de ruteo sobre cartografía OpenStreetMap"),
        ("process", "Cálculo de Ruta y Semáforo de Tráfico", "Segmentación cromática: Verde (Fluido), Amarillo (Medio), Rojo (Pesado)"),
        ("decision", "¿Usuario prefiere navegación asistida?", "Elección entre simulador web o aplicaciones nativas de GPS"),
        ("process", "Lanzamiento de Deep Links Externos", "Apertura de la ruta en la aplicación oficial de Waze o Google Maps"),
        ("start_end", "Navegación Activa hacia la Propiedad", "Acompañamiento en vivo durante el traslado")
    ], "4.8. Flujo del Motor de Navegación Vial y Waze (OSRM Engine)")

    # ==========================
    # 6. DICCIONARIO DE DATOS
    # ==========================
    doc.add_page_break()
    add_custom_heading("5. Estructura y Diccionario de Datos NoSQL (Cloud Firestore)", level=1, color=COLOR_TERRACOTA)

    fs_dict = [
        ["users", "uid (PK)", "String", "Identificador único emitido por Firebase Auth."],
        ["users", "email", "String", "Correo electrónico validado del usuario."],
        ["users", "role", "String", "Rol de control de acceso: 'seeker' o 'owner'."],
        ["users", "favorites", "Array<String>", "Lista de IDs de propiedades marcadas."],
        ["properties", "id (PK)", "String", "Identificador documental del anuncio."],
        ["properties", "ownerId", "String", "UID del propietario que creó el anuncio."],
        ["properties", "title", "String", "Título comercial del inmueble."],
        ["properties", "price", "Number", "Canon mensual o precio en moneda local."],
        ["properties", "location", "Map", "Coordenadas 'lat', 'lng' y dirección en texto."],
        ["properties", "images", "Array<String>", "URLs públicas seguras en Firebase Storage."],
        ["conversations", "participants", "Array<String>", "Arreglo con los 2 UIDs involucrados en el chat."],
        ["conversations", "lastMessage", "String", "Previsualización del último texto enviado."],
        ["messages", "senderId", "String", "UID del autor del mensaje emitido."],
        ["messages", "text", "String", "Contenido textual transmitido."],
        ["messages", "timestamp", "Timestamp", "Marca temporal de persistencia en servidor."],
        ["system_updates", "active", "Boolean", "Bandera de vigencia del aviso global."]
    ]

    tbl_fs = doc.add_table(rows=1, cols=4)
    tbl_fs.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fs = tbl_fs.rows[0].cells
    for i, h in enumerate(["Colección", "Campo", "Tipo de Dato", "Descripción del Negocio"]):
        c_fs[i].text = h
        set_cell_style(c_fs[i], bg_color="EA580C")
        c_fs[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c_fs[i].paragraphs[0].runs[0].font.bold = True
        c_fs[i].paragraphs[0].runs[0].font.name = "Arial"
        c_fs[i].paragraphs[0].runs[0].font.size = Pt(9)

    for row in fs_dict:
        r_c = tbl_fs.add_row().cells
        for idx, val in enumerate(row):
            r_c[idx].text = val
            set_cell_style(r_c[idx], bg_color="FFFFFF" if idx != 0 else "F8FAFC")
            r_c[idx].paragraphs[0].runs[0].font.name = "Arial"
            r_c[idx].paragraphs[0].runs[0].font.size = Pt(8.5)

    # ==========================
    # 7. TRAZABILIDAD Y PRUEBAS
    # ==========================
    doc.add_page_break()
    add_custom_heading("6. Trazabilidad de Requerimientos y Casos de Prueba (QA)", level=1, color=COLOR_TERRACOTA)

    add_custom_heading("6.1. Matriz de Trazabilidad de Requerimientos (RTM)", level=2, color=COLOR_CAFE)
    rtm_data = [
        ["RF-01", "Autenticación y Roles", "AuthContext.jsx / LoginScreen.kt", "TC-AUTH-01", "Aprobado"],
        ["RF-02", "Control Inactividad (5 min)", "RuwaJay Security Modal / LocalStorage", "TC-SEC-01", "Aprobado"],
        ["RF-03", "Reset de Contraseña", "ResetPasswordPage.jsx (oobCode)", "TC-AUTH-02", "Aprobado"],
        ["RF-04", "Publicación de Inmuebles", "PublishPage.jsx / Storage rules", "TC-PROP-01", "Aprobado"],
        ["RF-05", "Gestión de Favoritos", "FavoritesContext.jsx / ProfilesScreen.kt", "TC-FAV-01", "Aprobado"],
        ["RF-06", "Comparador (Máx 3)", "PropertyCompareModal.jsx / CompareContext", "TC-COMP-01", "Aprobado"],
        ["RF-07", "Chat en Tiempo Real", "ChatPage.jsx / conversations listener", "TC-CHAT-01", "Aprobado"],
        ["RF-08", "Navegación OSRM / Waze", "RoutePage.jsx / Waze deep link", "TC-NAV-01", "Aprobado"]
    ]

    tbl_rtm = doc.add_table(rows=1, cols=5)
    tbl_rtm.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_rt = tbl_rtm.rows[0].cells
    for i, h in enumerate(["ID Req", "Requerimiento", "Módulo / Componente", "ID Prueba", "Estado QA"]):
        c_rt[i].text = h
        set_cell_style(c_rt[i], bg_color="1E293B")
        c_rt[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c_rt[i].paragraphs[0].runs[0].font.bold = True
        c_rt[i].paragraphs[0].runs[0].font.name = "Arial"
        c_rt[i].paragraphs[0].runs[0].font.size = Pt(8.5)

    for row in rtm_data:
        r_c = tbl_rtm.add_row().cells
        for idx, val in enumerate(row):
            r_c[idx].text = val
            set_cell_style(r_c[idx], bg_color="F0FDF4" if idx == 4 else ("FFFFFF" if idx != 0 else "F8FAFC"))
            r_c[idx].paragraphs[0].runs[0].font.name = "Arial"
            r_c[idx].paragraphs[0].runs[0].font.size = Pt(8)

    doc.add_paragraph()
    add_custom_heading("6.2. Casos de Prueba Ejecutados (Test Cases)", level=2, color=COLOR_CAFE)
    tc_data = [
        ["TC-SEC-01", "Inactividad multi-pestaña", "Dejar la estación desatendida 5 min sin eventos DOM.", "Modal de advertencia con cuenta de 60s y posterior logout forzoso.", "Pass"],
        ["TC-AUTH-02", "Intercepción de oobCode", "Hacer clic en enlace de correo /reset-password.", "Captura de token en React y cambio de clave sin tocar pantalla blanca.", "Pass"],
        ["TC-COMP-01", "Límite del comparador", "Intentar agregar una 4ta propiedad al comparador.", "Bloqueo de inserción y despliegue de alerta toast informativa.", "Pass"],
        ["TC-NAV-01", "Fallo de red en OSRM", "Simular pérdida de conectividad durante cálculo de ruta.", "Activación de fallback ortodrómico y opción directa a app Waze.", "Pass"]
    ]

    tbl_tc = doc.add_table(rows=1, cols=5)
    tbl_tc.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_tc = tbl_tc.rows[0].cells
    for i, h in enumerate(["ID Prueba", "Escenario", "Acción de Entrada", "Resultado Esperado", "Veredicto"]):
        c_tc[i].text = h
        set_cell_style(c_tc[i], bg_color="EA580C")
        c_tc[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c_tc[i].paragraphs[0].runs[0].font.bold = True
        c_tc[i].paragraphs[0].runs[0].font.name = "Arial"
        c_tc[i].paragraphs[0].runs[0].font.size = Pt(8.5)

    for row in tc_data:
        r_c = tbl_tc.add_row().cells
        for idx, val in enumerate(row):
            r_c[idx].text = val
            set_cell_style(r_c[idx], bg_color="F0FDF4" if idx == 4 else ("FFFFFF" if idx != 0 else "F8FAFC"))
            r_c[idx].paragraphs[0].runs[0].font.name = "Arial"
            r_c[idx].paragraphs[0].runs[0].font.size = Pt(8)

    # ==========================
    # 8. GUARDADO FINAL
    # ==========================
    output_filename = "Documentacion_Tecnica_RuwaJay_Ingenieria.docx"
    doc.save(output_filename)
    print(f"\n[ÉXITO] Documento generado correctamente: {output_filename}")

if __name__ == "__main__":
    build_engineering_document()